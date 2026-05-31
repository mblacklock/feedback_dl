import io
import json

import openpyxl
from django.test import TestCase
from django.urls import reverse

from .utils import build_marking_workbook, parse_builder_payload, validate_builder_config


class MarkingSheetBuilderTests(TestCase):
    def test_builder_get_renders_page(self):
        response = self.client.get(reverse("marking_sheet_builder"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Marking Sheet Builder")
        self.assertContains(response, "Rubric mark")

    def test_payload_defaults_to_valid_template(self):
        config = parse_builder_payload("{}")

        self.assertEqual(config["name_mode"], "full")
        self.assertEqual(config["degree_level"], "BEng")
        self.assertEqual(config["rows"], 100)
        self.assertEqual(config["max_mark"], 100)
        self.assertEqual(config["columns"][0]["type"], "numeric")
        self.assertEqual(config["columns"][0]["max_mark"], 100)
        self.assertIn("First", config["rubric_boundaries"])

    def test_download_workbook_has_headers_and_rubric_dropdown(self):
        payload = {
            "name_mode": "split",
            "rows": 5,
            "max_mark": 100,
            "degree_level": "BEng",
            "columns": [
                {"type": "numeric", "title": "Coursework Mark", "max_mark": 50},
                {
                    "type": "rubric",
                    "title": "Rubric Grade",
                    "max_mark": 50,
                    "subdivision": "high_low",
                },
                {"type": "information", "title": "Submission Notes"},
                {"type": "feedback", "title": "Feedback"},
            ],
        }

        response = self.client.post(
            reverse("marking_sheet_builder"),
            {"builder_payload": json.dumps(payload)},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn("marking_sheet_template.xlsx", response["Content-Disposition"])

        workbook = openpyxl.load_workbook(io.BytesIO(response.content))
        sheet = workbook["Marking Sheet"]
        self.assertEqual(
            [cell.value for cell in sheet[1]],
            [
                "Student ID",
                "Last Name",
                "First Name",
                "Coursework Mark (50)",
                "Rubric Grade (50)",
                "Submission Notes",
                "Mark (100)",
                "%",
                "Feedback",
            ],
        )
        self.assertEqual(workbook["Rubric Boundaries"]["A1"].value, "Rubric Grade")
        self.assertEqual(workbook["Rubric Boundaries"]["B1"].value, "Rubric Grade marks")
        self.assertEqual(workbook["Rubric Boundaries"]["A2"].value, "Max 1st")
        self.assertEqual(workbook["Rubric Boundaries"]["B2"].value, 50)
        self.assertIn("High 2:1", [cell.value for cell in workbook["Rubric Boundaries"]["A"]])
        self.assertEqual(workbook["Rubric Boundaries"].sheet_state, "hidden")
        self.assertIsNone(sheet.auto_filter.ref)
        self.assertTrue(sheet["G2"].value.startswith("=IF(OR("))
        self.assertIn("VLOOKUP", sheet["G2"].value)
        self.assertIn("N(D2)", sheet["G2"].value)
        self.assertEqual(sheet["H2"].value, '=IF(G2<>"",G2/100,"")')
        self.assertEqual(sheet["H2"].number_format, "0%")

        validations = list(sheet.data_validations.dataValidation)
        self.assertTrue(any(validation.type == "list" for validation in validations))
        self.assertTrue(any(validation.type == "decimal" for validation in validations))
        decimal_validation = next(validation for validation in validations if validation.type == "decimal")
        self.assertEqual(decimal_validation.formula2, "50")

    def test_build_workbook_supports_full_name_mode(self):
        config = parse_builder_payload(
            json.dumps(
                {
                    "name_mode": "full",
                    "columns": [{"type": "feedback", "title": "Tutor Comment"}],
                }
            )
        )

        workbook_bytes = build_marking_workbook(config)
        workbook = openpyxl.load_workbook(workbook_bytes)
        headers = [cell.value for cell in workbook["Marking Sheet"][1]]

        self.assertEqual(headers[:2], ["Student ID", "Student Name"])
        self.assertIn("Tutor Comment", headers)
        self.assertEqual(headers[-3:], ["Mark (100)", "%", "Tutor Comment"])

    def test_numeric_column_maxima_must_sum_to_total_mark(self):
        config = parse_builder_payload(
            json.dumps(
                {
                    "max_mark": 100,
                    "columns": [
                        {"type": "numeric", "title": "Coursework", "max_mark": 40},
                        {"type": "numeric", "title": "Exam", "max_mark": 50},
                    ],
                }
            )
        )

        errors = validate_builder_config(config)

        self.assertEqual(
            errors,
            ["Mark and rubric column maxima must add up to 100. They currently add up to 90."],
        )

    def test_post_with_mismatched_numeric_total_renders_error(self):
        payload = {
            "max_mark": 100,
            "columns": [
                {"type": "numeric", "title": "Coursework", "max_mark": 40},
                {"type": "numeric", "title": "Exam", "max_mark": 50},
            ],
        }

        response = self.client.post(
            reverse("marking_sheet_builder"),
            {"builder_payload": json.dumps(payload)},
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Mark and rubric column maxima must add up to 100")
        self.assertNotEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    def test_rubric_column_maxima_count_towards_total_mark(self):
        config = parse_builder_payload(
            json.dumps(
                {
                    "max_mark": 100,
                    "degree_level": "MEng/MSc",
                    "columns": [
                        {"type": "numeric", "title": "Quiz", "max_mark": 20},
                        {
                            "type": "rubric",
                            "title": "Portfolio",
                            "max_mark": 80,
                            "subdivision": "high_mid_low",
                        },
                    ],
                }
            )
        )

        self.assertEqual(validate_builder_config(config), [])
        workbook_bytes = build_marking_workbook(config)
        workbook = openpyxl.load_workbook(workbook_bytes)
        rubric_labels = [cell.value for cell in workbook["Rubric Boundaries"]["A"]]

        self.assertIn("Max 1st/Dist", rubric_labels)
        self.assertIn("Mid 2:1/Merit", rubric_labels)
        self.assertNotIn("3rd", " ".join(str(label) for label in rubric_labels))
        self.assertIn("VLOOKUP", workbook["Marking Sheet"]["E2"].value)
        self.assertEqual(workbook["Marking Sheet"]["F2"].value, '=IF(E2<>"",E2/100,"")')

    def test_feedback_columns_placed_at_the_end(self):
        config = parse_builder_payload(
            json.dumps(
                {
                    "max_mark": 100,
                    "columns": [
                        {"type": "feedback", "title": "First Feedback"},
                        {"type": "numeric", "title": "Quiz", "max_mark": 50},
                        {"type": "feedback", "title": "Second Feedback"},
                        {"type": "numeric", "title": "Exam", "max_mark": 50},
                    ],
                }
            )
        )

        columns = config["columns"]
        self.assertEqual(columns[0]["type"], "numeric")
        self.assertEqual(columns[1]["type"], "numeric")
        self.assertEqual(columns[2]["type"], "feedback")
        self.assertEqual(columns[2]["title"], "First Feedback")
        self.assertEqual(columns[3]["type"], "feedback")
        self.assertEqual(columns[3]["title"], "Second Feedback")
