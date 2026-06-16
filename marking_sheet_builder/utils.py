import io
import json
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter, quote_sheetname
from openpyxl.worksheet.datavalidation import DataValidation

from core.utils.grade_bands import calculate_grade_bands


def get_custom_grade_bands(max_mark, subdivision, degree_level, custom_percentages=None, column_marks=None):
    from math import floor
    from core.utils.grade_bands import get_custom_grade_bands as core_get_custom_grade_bands
    
    if column_marks and isinstance(column_marks, dict):
        return core_get_custom_grade_bands(max_mark, subdivision, degree_level, custom_marks=column_marks)
        
    # Legacy percentage boundaries fallback
    default_bands = calculate_grade_bands(max_mark, subdivision, degree_level)
    if not custom_percentages:
        return default_bands
        
    sub_custom = custom_percentages.get(subdivision)
    if not sub_custom or not isinstance(sub_custom, dict):
        return default_bands
        
    new_bands = []
    for band in default_bands:
        grade_name = band["grade"]
        
        # Strip M-level suffixes to look up the correct base grade
        ug_grade_name = grade_name
        if degree_level and degree_level.strip().lower().startswith('m'):
            ug_grade_name = (
                ug_grade_name.replace('1st/Dist', '1st')
                             .replace('2:1/Merit', '2:1')
                             .replace('2:2/Pass', '2:2')
            )
            
        custom_pct = sub_custom.get(ug_grade_name)
        if custom_pct is not None:
            try:
                pct = float(custom_pct) / 100.0
                custom_mark = int(floor((max_mark * pct) + 0.5))
                new_bands.append({"grade": grade_name, "marks": custom_mark})
            except (ValueError, TypeError):
                new_bands.append(band)
        else:
            new_bands.append(band)
            
    return new_bands


ALLOWED_COLUMN_TYPES = {
    "numeric": "Numeric mark",
    "rubric": "Rubric mark",
    "information": "Information",
    "feedback": "Feedback comment",
}
ALLOWED_SUBDIVISIONS = {"none", "high_low", "high_mid_low"}
ALLOWED_DEGREE_LEVELS = {"BEng", "MEng/MSc"}


def parse_builder_payload(raw_payload):
    try:
        payload = json.loads(raw_payload or "{}")
    except json.JSONDecodeError:
        payload = {}

    name_mode = payload.get("name_mode")
    if name_mode not in {"full", "split"}:
        name_mode = "full"

    degree_level = payload.get("degree_level")
    if degree_level not in ALLOWED_DEGREE_LEVELS:
        degree_level = "BEng"

    rows = _parse_positive_int(payload.get("rows"), default=100, maximum=500)
    max_mark = _parse_positive_int(payload.get("max_mark"), default=100, maximum=10000)

    boundaries = _clean_lines(payload.get("rubric_boundaries"))
    if not boundaries:
        boundaries = ["First", "Upper second", "Lower second", "Third", "Fail"]

    columns = []
    for item in payload.get("columns", []):
        column_type = item.get("type")
        if column_type not in ALLOWED_COLUMN_TYPES:
            continue
        title = _clean_title(item.get("title")) or ALLOWED_COLUMN_TYPES[column_type]
        column = {"type": column_type, "title": title}
        if column_type in {"numeric", "rubric"}:
            column["max_mark"] = _parse_positive_int(
                item.get("max_mark"), default=max_mark, maximum=10000
            )
        if column_type == "rubric":
            subdivision = item.get("subdivision")
            if subdivision not in ALLOWED_SUBDIVISIONS:
                subdivision = "none"
            column["subdivision"] = subdivision
            
            # Parse custom integer marks
            marks = item.get("marks")
            if isinstance(marks, dict):
                clean_marks = {}
                for g, val in marks.items():
                    try:
                        clean_marks[str(g)] = int(val)
                    except (ValueError, TypeError):
                        pass
                column["marks"] = clean_marks
        columns.append(column)

    if not columns:
        columns = [{"type": "numeric", "title": "Mark", "max_mark": max_mark}]
    else:
        # Stable sort: Feedback comment columns should come at the end
        non_feedback = [c for c in columns if c["type"] != "feedback"]
        feedback = [c for c in columns if c["type"] == "feedback"]
        columns = non_feedback + feedback

    rubric_custom_percentages = payload.get("rubric_custom_percentages")
    if not isinstance(rubric_custom_percentages, dict):
        rubric_custom_percentages = {}

    return {
        "name_mode": name_mode,
        "degree_level": degree_level,
        "rows": rows,
        "max_mark": max_mark,
        "rubric_boundaries": boundaries,
        "columns": columns,
        "rubric_custom_percentages": rubric_custom_percentages,
    }


def validate_builder_config(config):
    assessed_total = sum(
        column.get("max_mark", 0)
        for column in config["columns"]
        if column["type"] in {"numeric", "rubric"}
    )
    if assessed_total != config["max_mark"]:
        return [
            f"Mark and rubric column maxima must add up to {config['max_mark']}. "
            f"They currently add up to {assessed_total}."
        ]
    return []


def build_marking_workbook(config):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Marking Sheet"

    rubric_sheet = workbook.create_sheet("Rubric Boundaries")
    _write_rubric_boundaries(rubric_sheet, config)
    rubric_sheet.sheet_state = "hidden"

    headers = ["Student ID"]
    if config["name_mode"] == "split":
        headers.extend(["Last Name", "First Name"])
    else:
        headers.append("Student Name")
    non_feedback_cols = [c for c in config["columns"] if c["type"] != "feedback"]
    feedback_cols = [c for c in config["columns"] if c["type"] == "feedback"]

    headers.extend(_display_column_title(column) for column in non_feedback_cols)
    headers.extend([f"Mark ({config['max_mark']})", "%"])
    headers.extend(_display_column_title(column) for column in feedback_cols)

    sheet.append(headers)
    data_rows = config["rows"]
    total_columns = len(headers)
    for _ in range(data_rows):
        sheet.append([""] * total_columns)

    _write_calculated_result_formulas(sheet, config, data_rows, total_columns)
    _style_marking_sheet(sheet, headers, data_rows, config)
    _apply_validations(sheet, config, data_rows)

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def _apply_validations(sheet, config, data_rows):
    first_marking_col = 3 if config["name_mode"] == "full" else 4
    non_feedback_cols = [c for c in config["columns"] if c["type"] != "feedback"]
    for idx, column in enumerate(non_feedback_cols):
        excel_col = get_column_letter(first_marking_col + idx)
        cell_range = f"{excel_col}2:{excel_col}{data_rows + 1}"
        if column["type"] == "numeric":
            validation = DataValidation(
                type="decimal",
                operator="between",
                formula1=0,
                formula2=column["max_mark"],
                allow_blank=True,
            )
            validation.error = f"Enter a number from 0 to {column['max_mark']}."
            validation.errorTitle = "Invalid mark"
            sheet.add_data_validation(validation)
            validation.add(cell_range)
        elif column["type"] == "rubric":
            rubric_boundaries = get_custom_grade_bands(
                column["max_mark"],
                column.get("subdivision", "none"),
                degree_level=config.get("degree_level", "BEng"),
                custom_percentages=config.get("rubric_custom_percentages"),
                column_marks=column.get("marks"),
            )
            original_offset = config["columns"].index(column)
            rubric_label_col, _ = _rubric_sheet_columns_for(config["columns"], original_offset)
            rubric_range = (
                f"{quote_sheetname('Rubric Boundaries')}!"
                f"${rubric_label_col}$2:${rubric_label_col}${len(rubric_boundaries) + 1}"
            )
            validation = DataValidation(
                type="list",
                formula1=rubric_range,
                allow_blank=True,
            )
            validation.error = "Choose one of the generated rubric boundaries."
            validation.errorTitle = "Invalid rubric boundary"
            sheet.add_data_validation(validation)
            validation.add(cell_range)


def _style_marking_sheet(sheet, headers, data_rows, config):
    primary_color = config.get("theme_primary", "#1a1a2e").lstrip("#")
    header_fill = PatternFill("solid", fgColor=primary_color)
    header_font = Font(bold=True, color="FFFFFF")
    thin = Side(style="thin", color="D9E2F3")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    for row in sheet.iter_rows(min_row=2, max_row=data_rows + 1, max_col=len(headers)):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    widths = {
        "Student ID": 16,
        "Student Name": 26,
        "Last Name": 20,
        "First Name": 20,
    }
    for index, header in enumerate(headers, start=1):
        width = widths.get(header, 24)
        if re.search(r"feedback|comment", header, re.IGNORECASE):
            width = 42
        sheet.column_dimensions[get_column_letter(index)].width = width

    sheet.freeze_panes = "A2"


def _write_calculated_result_formulas(sheet, config, data_rows, total_columns):
    first_marking_col = 3 if config["name_mode"] == "full" else 4
    non_feedback_cols = [c for c in config["columns"] if c["type"] != "feedback"]
    
    mark_col = get_column_letter(first_marking_col + len(non_feedback_cols))
    percent_col = get_column_letter(first_marking_col + len(non_feedback_cols) + 1)
    
    assessed_cols = [
        (non_feedback_cols.index(column), column)
        for column in config["columns"]
        if column["type"] in {"numeric", "rubric"}
    ]

    for row_index in range(2, data_rows + 2):
        terms = []
        presence_checks = []
        for idx, column in assessed_cols:
            source_col = get_column_letter(first_marking_col + idx)
            source_cell = f"{source_col}{row_index}"
            presence_checks.append(f'{source_cell}<>""')
            if column["type"] == "numeric":
                terms.append(f"N({source_cell})")
            elif column["type"] == "rubric":
                original_offset = config["columns"].index(column)
                label_col, marks_col = _rubric_sheet_columns_for(config["columns"], original_offset)
                band_count = len(
                    get_custom_grade_bands(
                        column["max_mark"],
                        column.get("subdivision", "none"),
                        degree_level=config.get("degree_level", "BEng"),
                        custom_percentages=config.get("rubric_custom_percentages"),
                        column_marks=column.get("marks"),
                    )
                )
                lookup_range = (
                    f"{quote_sheetname('Rubric Boundaries')}!"
                    f"${label_col}$2:${marks_col}${band_count + 1}"
                )
                terms.append(f"IFERROR(VLOOKUP({source_cell},{lookup_range},2,FALSE),0)")

        if not terms:
            sheet[f"{mark_col}{row_index}"] = '=""'
            sheet[f"{percent_col}{row_index}"] = '=""'
            continue

        formula = f'=IF(OR({",".join(presence_checks)}),{"+".join(terms)},"")'
        sheet[f"{mark_col}{row_index}"] = formula
        sheet[f"{percent_col}{row_index}"] = (
            f'=IF({mark_col}{row_index}<>"",{mark_col}{row_index}/{config["max_mark"]},"")'
        )
        sheet[f"{percent_col}{row_index}"].number_format = "0%"


def _display_column_title(column):
    if column["type"] in {"numeric", "rubric"}:
        return f"{column['title']} ({column['max_mark']})"
    return column["title"]


def _write_rubric_boundaries(sheet, config):
    rubric_columns = [
        column for column in config["columns"] if column["type"] == "rubric"
    ]
    if not rubric_columns:
        sheet["A1"] = "Rubric Boundary"
        for index, boundary in enumerate(config["rubric_boundaries"], start=2):
            sheet.cell(row=index, column=1, value=boundary)
        sheet.column_dimensions["A"].width = 28
        return

    for rubric_index, column in enumerate(rubric_columns):
        label_col_index = (rubric_index * 2) + 1
        marks_col_index = label_col_index + 1
        label_excel_col = get_column_letter(label_col_index)
        marks_excel_col = get_column_letter(marks_col_index)
        sheet.cell(row=1, column=label_col_index, value=column["title"])
        sheet.cell(row=1, column=marks_col_index, value=f"{column['title']} marks")
        primary_color = config.get("theme_primary", "#1a1a2e").lstrip("#")
        for header_cell in (
            sheet.cell(row=1, column=label_col_index),
            sheet.cell(row=1, column=marks_col_index),
        ):
            header_cell.font = Font(bold=True, color="FFFFFF")
            header_cell.fill = PatternFill("solid", fgColor=primary_color)

        bands = get_custom_grade_bands(
            column["max_mark"],
            column.get("subdivision", "none"),
            degree_level=config.get("degree_level", "BEng"),
            custom_percentages=config.get("rubric_custom_percentages"),
            column_marks=column.get("marks"),
        )
        for row_index, band in enumerate(bands, start=2):
            sheet.cell(row=row_index, column=label_col_index, value=band["grade"])
            sheet.cell(row=row_index, column=marks_col_index, value=band["marks"])
        sheet.column_dimensions[label_excel_col].width = 28
        sheet.column_dimensions[marks_excel_col].width = 14


def _rubric_sheet_columns_for(columns, marking_offset):
    rubric_index = 0
    for offset, column in enumerate(columns):
        if column["type"] != "rubric":
            continue
        if offset == marking_offset:
            label_col = get_column_letter((rubric_index * 2) + 1)
            marks_col = get_column_letter((rubric_index * 2) + 2)
            return label_col, marks_col
        rubric_index += 1
    return "A", "B"


def _clean_lines(value):
    if isinstance(value, list):
        candidates = value
    else:
        candidates = str(value or "").splitlines()
    return [_clean_title(item) for item in candidates if _clean_title(item)]


def _clean_title(value):
    return re.sub(r"\s+", " ", str(value or "")).strip()[:80]


def _parse_positive_int(value, default, maximum):
    try:
        parsed = int(value)
    except (TypeError, ValueError):
        return default
    return max(1, min(parsed, maximum))
