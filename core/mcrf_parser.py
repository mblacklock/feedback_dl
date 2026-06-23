import io
import re
import openpyxl
from openpyxl.utils import get_column_letter

try:
    import pypdf
except ImportError:
    pypdf = None

MODULE_CODE_RE = re.compile(r'^([A-Z]{2,6}\d{3,5}[A-Z]?)\s*[-:–]?\s*(.*)$')

def normalize_comp_code(code_str):
    if not code_str:
        return ""
    s = str(code_str).strip().split('.')[0]
    if s.isdigit():
        return f"{int(s):03d}"
    return s.upper()

def clean_component_title(title):
    if not title:
        return ""
    cleaned = re.sub(
        r'\s*\([^)]*\b(words?|hrs?|hours?|mins?|minutes?)\b[^)]*\)',
        '',
        title,
        flags=re.IGNORECASE
    )
    cleaned = re.sub(r'\s*\(\s*equivalent\s*\)', '', cleaned, flags=re.IGNORECASE)
    return cleaned.strip().rstrip(',').strip()

def is_assessment_component_column(header_name):
    """
    Detects if a column header represents an assessment component mark.
    Matches columns containing 'mark', 'score', 'result', 'pct', 'percent',
    or common assessment keywords/codes (e.g. CW1, Exam, Quiz, Project, Assignment),
    while excluding student info, grades (letter grades), and overall module totals/averages.
    """
    hl = str(header_name).lower()
    
    # Exclude student info, occurrences, dates, grades (letter/verbal grades), and overall totals
    if any(kw in hl for kw in [
        "name", "student", "id", "number", "username", "email",
        "total", "overall", "average", "final", "module", "weighted",
        "occ", "period", "level", "year", "credits", "date", "grade",
        "gpa", "feedback", "comment", "marker", "tutor", "row"
    ]):
        return False
        
    # Match standard component headers containing keywords
    if any(kw in hl for kw in ["mark", "score", "result", "pct", "percent", "pts", "points"]):
        return True
        
    # Match common assessment prefixes/words (e.g. CW, Coursework, Exam, Test, Quiz, Project, Presentation, Portfolio)
    if re.search(r'\b(?:cw|coursework|exam|test|quiz|proj|project|pres|presentation|port|portfolio|ass|assignment|viva|prac|practical|comp|component)\d*\b', hl):
        return True
        
    return False

def parse_mcrf_pdf(file_file):
    """
    Parses a PDF MCRF file using pypdf's layout-preserving text extraction.
    Reconstructs headers, data_rows, and module_info to match spreadsheet outputs.
    """
    if not pypdf:
        raise ValueError("pypdf library is required to parse PDF files. Install it using 'pip install pypdf'.")
        
    reader = pypdf.PdfReader(file_file)
    layout_text_pages = []
    for page in reader.pages:
        try:
            layout_text = page.extract_text(extraction_mode="layout")
        except Exception:
            layout_text = page.extract_text()
        layout_text_pages.append(layout_text)
        
    full_text = "\n".join(layout_text_pages)
    lines = [line for line in full_text.splitlines()]
    
    in_components = False
    comp_names_map = {}
    components_list = []
    module_lines = []
    occurrence_lines = []
    
    year = ""
    credits_val = 20
    period = ""
    level = 4
    occurrence = ""
    
    for idx, line in enumerate(lines):
        line_clean = line.strip()
        
        # 1. Component parsing
        if "Component" in line and "Weighting" in line:
            in_components = True
            continue
            
        if in_components:
            leading_spaces = len(line) - len(line.lstrip())
            if leading_spaces > 5:
                in_components = False
                continue
                
            m = re.match(r'^\s*([A-Za-z0-9]+)\s+(.+?)\s+(\d+%)', line)
            if m:
                code, desc, weight = m.groups()
                code_norm = normalize_comp_code(code)
                desc_clean = clean_component_title(desc)
                weight_val = int(weight.replace("%", ""))
                comp_names_map[code_norm] = desc_clean
                if not any(c["code"] == code_norm for c in components_list):
                    components_list.append({
                        "code": code_norm,
                        "column": f"{code_norm} - {weight_val}% - Mark",
                        "weight": weight_val
                    })
            else:
                if line_clean:
                    in_components = False
                    
        # 2. Metadata parsing
        if line.startswith("Module ") and not module_lines and "Marks Record Form" not in line and "(MCRF)" not in line:
            tutor_idx = line.find("Tutor")
            if tutor_idx != -1:
                module_lines.append(line[len("Module"):tutor_idx].strip())
            else:
                module_lines.append(line[len("Module"):].strip())
                
            j = idx + 1
            while j < len(lines) and lines[j].startswith(" " * 10) and not any(k in lines[j] for k in ["Year", "Period", "Occurrence", "Component", "Student ID"]):
                module_lines.append(lines[j].strip())
                j += 1
                
        if line.startswith("Year ") and not year:
            credits_idx = line.find("Credits")
            if credits_idx != -1:
                year = line[len("Year"):credits_idx].strip()
                credits_str = line[credits_idx + len("Credits"):].strip()
                try:
                    credits_val = int(float(credits_str))
                except ValueError:
                    credits_val = 20
            else:
                year = line[len("Year"):].strip()
                
        if line.startswith("Period ") and not period:
            level_idx = line.find("Level")
            if level_idx != -1:
                period = line[len("Period"):level_idx].strip()
                level_str = line[level_idx + len("Level"):].strip()
                try:
                    level = int(level_str)
                except ValueError:
                    level = 4
            else:
                period = line[len("Period"):].strip()
                
        if line.startswith("Occurrence ") and not occurrence_lines:
            loc_idx = line.find("Location")
            if loc_idx != -1:
                occurrence_lines.append(line[len("Occurrence"):loc_idx].strip())
            else:
                occurrence_lines.append(line[len("Occurrence"):].strip())
                
            j = idx + 1
            while j < len(lines) and lines[j].startswith(" " * 10) and not any(k in lines[j] for k in ["Component", "Student ID"]):
                occurrence_lines.append(lines[j].strip())
                j += 1
                
    module_full = " ".join(module_lines)
    m_match = MODULE_CODE_RE.match(module_full)
    if m_match:
        module_code = m_match.group(1).strip()
        module_title = m_match.group(2).strip()
    else:
        module_code = ""
        module_title = module_full
        
    occ_full = " ".join(occurrence_lines)
    occ_codes = re.findall(r'\b([A-Z]{1,4}NN)\b', occ_full)
    occurrence = "/".join(occ_codes) if occ_codes else occ_full
    
    if "/" in year:
        parts = year.split("/")
        if len(parts) == 2 and len(parts[1]) == 1:
            year = f"{parts[0]}/2{parts[1]}"
            
    # Detect dynamic offsets from the column header line
    header_line = ""
    for line in lines:
        if "Student ID" in line and "Occ" in line:
            header_line = line
            break
            
    use_offsets = False
    if header_line:
        occ_idx = header_line.find("Occ")
        period_idx = header_line.find("Period")
        mark_indices = [m.start() for m in re.finditer(r'\bMark\b', header_line)]
        grade_indices = [m.start() for m in re.finditer(r'\bGrade\b', header_line)]
        if occ_idx != -1 and period_idx != -1 and len(mark_indices) > 0 and len(grade_indices) > 0:
            use_offsets = True
            first_mark = mark_indices[0]
            period_end = period_idx + len("Period")
            start_search_idx = max(period_end + 2, first_mark - 20)

    student_id_re = re.compile(r'^\s*(\d{8}(?:/\d+)?)\s+(.*)$')
    students_list = []
    current_student = None
    
    for line in lines:
        m = student_id_re.match(line)
        if m:
            stud_id = m.group(1)
            
            if use_offsets:
                # Slicing based on dynamic header offsets and splitting text columns
                text_part = line[:start_search_idx].strip()
                parts = re.split(r'\s{2,}', text_part)
                
                name_start = parts[1] if len(parts) > 1 else ""
                occ = parts[2] if len(parts) > 2 else ""
                period_val = parts[3] if len(parts) > 3 else ""
                
                # Tokenize remaining part of the line
                tokens = []
                for tok_m in re.finditer(r'\S+', line[start_search_idx:]):
                    val = tok_m.group(0)
                    pos = tok_m.start() + start_search_idx
                    tokens.append((val, pos))
                
                col_targets = []
                for i in range(len(mark_indices)):
                    col_targets.append(mark_indices[i])
                    col_targets.append(grade_indices[i])
                
                def is_numeric(v):
                    try:
                        float(v)
                        return True
                    except ValueError:
                        return False
                
                # Search for the best shift from -15 to +15
                best_shift = 0
                min_error = float('inf')
                for shift in range(-15, 16):
                    error = 0
                    for val, pos in tokens:
                        col_idx = min(range(len(col_targets)), key=lambda idx: abs(pos - shift - col_targets[idx]))
                        target_pos = col_targets[col_idx]
                        dist = abs(pos - shift - target_pos)
                        
                        is_num = is_numeric(val)
                        is_mark_col = (col_idx % 2 == 0)
                        
                        penalty = 0
                        if is_num and not is_mark_col:
                            penalty = 1000
                        elif not is_num and is_mark_col:
                            penalty = 1000
                            
                        error += dist + penalty
                        
                    if error < min_error:
                        min_error = error
                        best_shift = shift
                        
                scores_grades = [""] * (2 * len(mark_indices))
                for val, pos in tokens:
                    col_idx = min(range(len(col_targets)), key=lambda idx: abs(pos - best_shift - col_targets[idx]))
                    if col_idx < len(scores_grades):
                        scores_grades[col_idx] = val
            else:
                rest = m.group(2)
                parts = re.split(r'\s{2,}', rest.strip())
                name_start = parts[0]
                occ = parts[1]
                period_val = parts[2]
                scores_grades = parts[3:]
            
            current_student = {
                "Student ID": stud_id,
                "Student Name": name_start,
                "Occ": occ,
                "Period": period_val,
                "scores_grades": scores_grades
            }
            students_list.append(current_student)
        elif current_student is not None:
            line_strip = line.strip()
            if line_strip and not re.search(r'\s{2,}', line_strip) and line.startswith(" " * 10):
                current_student["Student Name"] += " " + line_strip
                
    headers = ["Student ID", "Student Name"]
    for comp in components_list:
        headers.append(f"{comp['code']} - {comp['weight']}% - Mark")
        headers.append(f"{comp['code']} - {comp['weight']}% - Grade")
    headers.extend(["Module - Mark", "Module - Grade"])
    
    data_rows = []
    for s in students_list:
        row = {
            "Student ID": s["Student ID"],
            "Student Name": s["Student Name"],
            "Occ": s["Occ"],
            "Period": s["Period"]
        }
        sg = s["scores_grades"]
        idx = 0
        for comp in components_list:
            mark_val = sg[idx] if idx < len(sg) else None
            grade_val = sg[idx+1] if idx+1 < len(sg) else None
            
            try:
                if mark_val is not None:
                    if mark_val.isdigit():
                        mark_val = int(mark_val)
                    elif mark_val == "":
                        mark_val = None
            except ValueError:
                pass
                
            row[f"{comp['code']} - {comp['weight']}% - Mark"] = mark_val
            row[f"{comp['code']} - {comp['weight']}% - Grade"] = grade_val
            idx += 2
            
        module_mark = sg[idx] if idx < len(sg) else None
        module_grade = sg[idx+1] if idx+1 < len(sg) else None
        try:
            if module_mark is not None:
                if module_mark.isdigit():
                    module_mark = int(module_mark)
                elif module_mark == "":
                    module_mark = None
        except ValueError:
            pass
        row["Module - Mark"] = module_mark
        row["Module - Grade"] = module_grade
        
        data_rows.append(row)
        
    return headers, data_rows, True, {
        "module_code": module_code,
        "module_title": module_title,
        "comp_names_map": comp_names_map,
        "year": year,
        "period": period,
        "occurrence": occurrence,
        "credits": credits_val,
        "detected_credits": credits_val,
    }

def parse_mcrf_workbook(file_file):
    """
    Parses a spreadsheet in-memory. Detects MCRF signature block
    and dynamically merges multi-row headings for clarity (e.g. "CW1" + "Mark" -> "CW1 - Mark").
    Supports modern .xlsx (via openpyxl) and legacy .xls (via xlrd).
    """
    file_bytes = file_file.read()
    file_file.seek(0)
    
    if file_bytes.startswith(b'%PDF'):
        return parse_mcrf_pdf(io.BytesIO(file_bytes))
    
    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
    except Exception:
        try:
            import xlrd
            wb = xlrd.open_workbook(file_contents=file_bytes)
            ws = wb.sheet_by_index(0)
            rows = [ws.row_values(i) for i in range(ws.nrows)]
        except Exception as e:
            raise ValueError(f"Spreadsheet format not recognized or corrupted: {str(e)}")
            
    if not rows or len(rows) < 2:
        raise ValueError("The uploaded spreadsheet is empty.")
        
    is_mcrf = False
    header_row_idx = 0
    
    # Scan first 50 rows to detect MCRF signatures and find the main header row
    for i in range(min(len(rows), 50)):
        row_cells = [str(c).strip() if c is not None else "" for c in rows[i]]
        row_text = " ".join(row_cells).lower()
        if "module marks record form" in row_text or "(mcrf)" in row_text:
            is_mcrf = True
            
        if any(keyword in cell.lower() for cell in row_cells for keyword in ["student id", "student no", "student number", "username"]):
            header_row_idx = i
            break
            
    raw_headers = [str(c).strip() if c is not None else "" for c in rows[header_row_idx]]
    headers = list(raw_headers)
    
    # Replace empty headers with '<Column X>'
    for idx in range(len(headers)):
        if not headers[idx]:
            headers[idx] = f"<Column {get_column_letter(idx + 1)}>"
            
    # If MCRF, merge parent and child headers
    if is_mcrf and header_row_idx > 0:
        parent_row = [str(c).strip() if c is not None else "" for c in rows[header_row_idx - 1]]
        current_parent = ""
        for c in range(len(headers)):
            p_val = parent_row[c] if c < len(parent_row) else ""
            if p_val != "":
                current_parent = p_val
                
            m_val = raw_headers[c] # Use raw header to verify original contents
            m_val_lower = m_val.lower()
            if m_val != "" and current_parent != "" and current_parent != m_val:
                if any(k in m_val_lower for k in ["mark", "grad", "result", "score"]):
                    headers[c] = f"{current_parent} - {m_val}"
                    
    # Read rows
    data_rows = []
    for r in rows[header_row_idx + 1:]:
        if any(cell is not None for cell in r):
            row_dict = {}
            for idx, cell in enumerate(r):
                if idx < len(headers):
                    h_name = headers[idx]
                    if h_name:
                        row_dict[h_name] = cell
            data_rows.append(row_dict)

    # --- Module details detection ---
    # Try C7 first (row 6, col 2) — the standard MCRF cell for module info.
    # If that's empty, scan pre-header rows for any cell matching a module code pattern.
    module_code = ""
    module_title = ""

    def _try_parse_module_cell(val):
        """Return (code, title) if val looks like a module details cell, else ('', '')."""
        if not val:
            return "", ""
        text = str(val).strip()
        m = MODULE_CODE_RE.match(text)
        if m:
            return m.group(1).strip(), m.group(2).strip()
        return "", ""

    # 1. Try canonical MCRF position: row 7 (index 6), column C (index 2)
    if len(rows) > 6:
        c7_val = rows[6][2] if len(rows[6]) > 2 else None
        module_code, module_title = _try_parse_module_cell(c7_val)

    # 2. If not found, scan all pre-header rows for a matching cell
    if not module_code:
        for row in rows[:header_row_idx]:
            for cell in row:
                code, title = _try_parse_module_cell(cell)
                if code:
                    module_code, module_title = code, title
                    break
            if module_code:
                break

    # --- Assessment components mapping and descriptions scan ---
    comp_names_map = {}



    # 1. Locate the row where Column A says "Component"
    component_row_idx = -1
    for idx, row in enumerate(rows[:header_row_idx]):
        if row and row[0] and str(row[0]).strip().lower() == "component":
            component_row_idx = idx
            break



    # 2. Extract component codes (Col A) and titles (Col B) from rows directly below it
    if component_row_idx != -1:
        for row in rows[component_row_idx + 1:header_row_idx]:
            if row and row[0] is not None:
                code_val = str(row[0]).strip()
                if not code_val:
                    break
                code_norm = normalize_comp_code(code_val)
                if len(row) > 1 and row[1]:
                    title_clean = clean_component_title(str(row[1]).strip())
                    comp_names_map[code_norm] = title_clean

    # --- Year, Period, Occurrence, Credits ---
    # Standard MCRF layout: Year=row8/colC, Period=row9/colC, Occurrence=row10/colC (0-indexed: rows 7,8,9)
    # Credits=row8/colE (0-indexed: row 7, col 4)
    year = ""
    period = ""
    occurrence = ""
    credits_val = 20

    def _cell(row_idx, col_idx):
        if len(rows) > row_idx and len(rows[row_idx]) > col_idx:
            v = rows[row_idx][col_idx]
            return str(v).strip() if v is not None else ""
        return ""

    # Try canonical positions first
    year_raw = _cell(7, 2)
    period_raw = _cell(8, 2)
    occurrence_raw = _cell(9, 2)
    credits_raw = _cell(7, 4)

    # Validate by checking the labels in col A
    if str(rows[7][0]).strip().lower() == "year" if len(rows) > 7 and rows[7] else False:
        year = year_raw
    if len(rows) > 7 and len(rows[7]) > 3 and str(rows[7][3]).strip().lower() == "credits":
        try:
            credits_val = int(float(credits_raw)) if credits_raw else 20
        except ValueError:
            credits_val = 20

    if str(rows[8][0]).strip().lower() == "period" if len(rows) > 8 and rows[8] else False:
        period = period_raw
    if str(rows[9][0]).strip().lower() == "occurrence" if len(rows) > 9 and rows[9] else False:
        occurrence_raw_val = occurrence_raw
        # Extract just BNN/FNN code from e.g. "BNN: September start - Newcastle upon Tyne"
        occ_match = re.match(r'^([A-Z]{1,4}NN)\b', occurrence_raw_val, re.IGNORECASE)
        occurrence = occ_match.group(1).upper() if occ_match else occurrence_raw_val

    # Fallback: scan pre-header rows for labelled cells if not found
    for row in rows[:header_row_idx]:
        if not row:
            continue
        for col_idx, cell in enumerate(row):
            if not cell:
                continue
            label = str(cell).strip().lower()
            val = str(row[col_idx + 1]).strip() if col_idx + 1 < len(row) and row[col_idx + 1] is not None else ""
            if label == "year" and not year:
                year = val
            elif label == "period" and not period:
                period = val
            elif label == "occurrence" and not occurrence:
                occ_match = re.match(r'^([A-Z]{1,4}NN)\b', val, re.IGNORECASE)
                occurrence = occ_match.group(1).upper() if occ_match else val
            elif label == "credits":
                if val:
                    try:
                        val_clean = re.sub(r'[^\d.]', '', val)
                        credits_val = int(float(val_clean))
                    except ValueError:
                        pass

    return headers, data_rows, is_mcrf, {
        "module_code": module_code,
        "module_title": module_title,
        "comp_names_map": comp_names_map,
        "year": year,
        "period": period,
        "occurrence": occurrence,
        "credits": credits_val,
        "detected_credits": credits_val,
    }
