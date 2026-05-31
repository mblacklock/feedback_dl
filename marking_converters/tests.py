import io
import openpyxl
import base64
import unittest.mock
from django.test import TestCase
from django.urls import reverse
from django.core.files.uploadedfile import SimpleUploadedFile
from marking_converters.views import normalize_id

class MarkingConvertersTests(TestCase):
    def setUp(self):
        # Local marks CSV
        self.marks_csv_bytes = b"Student ID,Student Name,CW1,Exam\nw12345678,Alice Smith,85,90\n12345679/2,Bob Jones,75,80\n"
        
        # Blackboard CSV
        self.bb_csv_bytes = b"Username,Last Name,First Name,CW1_Col,Exam_Col\nw12345678,Smith,Alice,,\n12345679,Jones,Bob,,\n"

        # Mock MCRF XLSX Template
        self.mcrf_wb = openpyxl.Workbook()
        ws = self.mcrf_wb.active
        ws.append(["Module Marks Record Form (MCRF)"])
        ws.append(["COMP3002 - Advanced Engineering Software"])
        ws.append(["", "", "CW1", "CW1", "Exam", "Exam"])
        ws.append(["Student ID", "Student Name", "Mark", "Grade", "Mark", "Grade"])
        ws.append(["w12345678", "Alice Smith", "", "", "", ""])
        ws.append(["12345679/2", "Bob Jones", "", "", "", ""])
        
        self.mcrf_bytes = io.BytesIO()
        self.mcrf_wb.save(self.mcrf_bytes)
        self.mcrf_bytes.seek(0)

    def test_normalize_id(self):
        """Verify student ID normalization extracts raw 8-digit sequence."""
        self.assertEqual(normalize_id("w12345678"), "12345678")
        self.assertEqual(normalize_id("12345678/1"), "12345678")
        self.assertEqual(normalize_id("   12345678   "), "12345678")
        self.assertEqual(normalize_id("bob"), "bob")

    def test_dashboard_view(self):
        """GET /mcrf-converter/ returns 200 and loads landing page."""
        url = reverse("marking_converters")
        resp = self.client.get(url)
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "Marking Sheet Converters")
        self.assertContains(resp, "Gradebook Populator")
        self.assertContains(resp, "MCRF Populator")

    def test_merge_upload_flow(self):
        """POST /mcrf-converter/merge/upload/ parses uploads and redirects to match."""
        url = reverse("gradebook_upload")
        
        marks_file = SimpleUploadedFile("marks.csv", self.marks_csv_bytes, content_type="text/csv")
        bb_file = SimpleUploadedFile("bb.csv", self.bb_csv_bytes, content_type="text/csv")
        
        resp = self.client.post(url, {
            "marks_file": marks_file,
            "bb_file": bb_file
        })
        self.assertEqual(resp.status_code, 302)
        self.assertEqual(resp.url, reverse("gradebook_match"))
        
        # Verify sessions
        session = self.client.session
        self.assertIn("converter_local_headers", session)
        self.assertIn("converter_bb_headers", session)
        self.assertEqual(session["converter_local_headers"], ["Student ID", "Student Name", "CW1", "Exam"])

    def test_merge_match_and_download(self):
        """POST /mcrf-converter/merge/match/ merges files and returns CSV response."""
        session = self.client.session
        session["converter_local_headers"] = ["Student ID", "Student Name", "CW1", "Exam"]
        session["converter_local_rows"] = [
            {"Student ID": "w12345678", "Student Name": "Alice Smith", "CW1": "85", "Exam": "90"},
            {"Student ID": "12345679/2", "Student Name": "Bob Jones", "CW1": "75", "Exam": "80"}
        ]
        session["converter_bb_headers"] = ["Username", "Last Name", "First Name", "CW1_Col", "Exam_Col"]
        session["converter_bb_rows"] = [
            {"Username": "w12345678", "Last Name": "Smith", "First Name": "Alice", "CW1_Col": "", "Exam_Col": ""},
            {"Username": "12345679", "Last Name": "Jones", "First Name": "Bob", "CW1_Col": "", "Exam_Col": ""}
        ]
        session.save()

        url = reverse("gradebook_match")
        resp = self.client.post(url, {
            "local_id_col": "Student ID",
            "bb_id_col": "Username",
            "map_local_0": "CW1",
            "map_bb_0": "CW1_Col",
            "map_local_1": "Exam",
            "map_bb_1": "Exam_Col"
        })
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp["Content-Type"], "text/csv; charset=utf-8")
        self.assertIn("attachment; filename=", resp["Content-Disposition"])
        
        content = resp.content.decode("utf-8")
        self.assertIn("85", content)
        self.assertIn("90", content)
        self.assertIn("75", content)
        self.assertIn("80", content)

    def test_populate_upload_flow(self):
        """POST /mcrf-converter/populate/upload/ parses uploads and redirects."""
        url = reverse("mcrf_upload")
        
        marks_file = SimpleUploadedFile("marks.csv", self.marks_csv_bytes, content_type="text/csv")
        mcrf_file = SimpleUploadedFile("mcrf.xlsx", self.mcrf_bytes.read(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        
        resp = self.client.post(url, {
            "marks_file": marks_file,
            "mcrf_file": mcrf_file
        })
        self.assertEqual(resp.status_code, 302)
        self.assertEqual(resp.url, reverse("mcrf_match"))
        
    def test_populate_match_only_displays_mark_columns(self):
        """Verify that only columns with 'mark' in their name are passed as target mapping options in mcrf_mark_headers."""
        session = self.client.session
        session["converter_local_headers"] = ["Student ID", "CW1"]
        session["converter_local_rows"] = []
        session["converter_mcrf_headers"] = ["Student ID", "CW1 - Mark", "CW1 - Grade", "Exam - Mark", "Exam - Grade"]
        session["converter_mcrf_rows"] = []
        session["converter_mcrf_template"] = base64.b64encode(self.mcrf_bytes.getvalue()).decode("utf-8")
        session["converter_mcrf_name"] = "blank_MCRF.xlsx"
        session.save()
        
        url = reverse("mcrf_match")
        resp = self.client.get(url)
        self.assertEqual(resp.status_code, 200)
        self.assertIn("CW1 - Mark", resp.context["mcrf_mark_headers"])
        self.assertIn("Exam - Mark", resp.context["mcrf_mark_headers"])
        self.assertNotIn("CW1 - Grade", resp.context["mcrf_mark_headers"])
        self.assertNotIn("Exam - Grade", resp.context["mcrf_mark_headers"])

    def test_populate_match_and_download(self):
        """POST /mcrf-converter/populate/match/ populates MCRF template and returns XLSX response."""
        session = self.client.session
        session["converter_local_headers"] = ["Student ID", "Student Name", "CW1", "Exam"]
        session["converter_local_rows"] = [
            {"Student ID": "w12345678", "Student Name": "Alice Smith", "CW1": "85", "Exam": "90"},
            {"Student ID": "12345679/2", "Student Name": "Bob Jones", "CW1": "75", "Exam": "80"}
        ]
        session["converter_mcrf_headers"] = ["Student ID", "Student Name", "CW1 - Mark", "CW1 - Grade", "Exam - Mark", "Exam - Grade"]
        session["converter_mcrf_rows"] = [
            {"Student ID": "w12345678", "Student Name": "Alice Smith", "CW1 - Mark": "", "CW1 - Grade": "", "Exam - Mark": "", "Exam - Grade": ""},
            {"Student ID": "12345679/2", "Student Name": "Bob Jones", "CW1 - Mark": "", "CW1 - Grade": "", "Exam - Mark": "", "Exam - Grade": ""}
        ]
        session["converter_mcrf_template"] = base64.b64encode(self.mcrf_bytes.getvalue()).decode("utf-8")
        session["converter_mcrf_name"] = "blank_MCRF.xlsx"
        session.save()

        url = reverse("mcrf_match")
        resp = self.client.post(url, {
            "local_id_col": "Student ID",
            "mcrf_id_col": "Student ID",
            "map_local_0": "CW1",
            "map_mcrf_0": "CW1 - Mark",
            "map_local_1": "Exam",
            "map_mcrf_1": "Exam - Mark"
        })
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.assertIn("attachment; filename=", resp["Content-Disposition"])
        
        # Verify workbook content
        wb = openpyxl.load_workbook(io.BytesIO(resp.content))
        ws = wb.active
        self.assertEqual(ws.cell(row=5, column=3).value, 85)
        self.assertEqual(ws.cell(row=5, column=5).value, 90)
        self.assertEqual(ws.cell(row=6, column=3).value, 75)
        self.assertEqual(ws.cell(row=6, column=5).value, 80)

    def test_populate_match_and_download_xls(self):
        """Verify MCRF Populator seamlessly populates legacy XLS templates and preserves styles."""
        import xlwt
        import xlrd
        
        # Create a real styled legacy XLS workbook in-memory
        xls_wb = xlwt.Workbook()
        ws = xls_wb.add_sheet("Sheet1")
        
        ws.write(0, 0, "Module Marks Record Form (MCRF)")
        ws.write(1, 0, "")
        ws.write(2, 2, "CW1")
        ws.write(2, 3, "CW1")
        ws.write(2, 4, "Exam")
        ws.write(2, 5, "Exam")
        
        ws.write(3, 0, "Student ID")
        ws.write(3, 1, "Student Name")
        ws.write(3, 2, "Mark")
        ws.write(3, 3, "Grade")
        ws.write(3, 4, "Mark")
        ws.write(3, 5, "Grade")
        
        # Bold and border styling to verify style preservation
        custom_style = xlwt.easyxf('font: bold on; border: left thin, right thin, top thin, bottom thin')
        
        ws.write(4, 0, "w12345678")
        ws.write(4, 1, "Alice Smith")
        ws.write(4, 2, "", custom_style)
        ws.write(4, 3, "")
        ws.write(4, 4, "", custom_style)
        ws.write(4, 5, "")
        
        ws.write(5, 0, "12345679/2")
        ws.write(5, 1, "Bob Jones")
        ws.write(5, 2, "", custom_style)
        ws.write(5, 3, "")
        ws.write(5, 4, "", custom_style)
        ws.write(5, 5, "")
        
        xls_io = io.BytesIO()
        xls_wb.save(xls_io)
        xls_bytes = xls_io.getvalue()
        
        session = self.client.session
        session["converter_local_headers"] = ["Student ID", "Student Name", "CW1", "Exam"]
        session["converter_local_rows"] = [
            {"Student ID": "w12345678", "Student Name": "Alice Smith", "CW1": "85", "Exam": "90"},
            {"Student ID": "12345679/2", "Student Name": "Bob Jones", "CW1": "75", "Exam": "80"}
        ]
        session["converter_mcrf_headers"] = ["Student ID", "Student Name", "CW1 - Mark", "CW1 - Grade", "Exam - Mark", "Exam - Grade"]
        session["converter_mcrf_rows"] = [
            {"Student ID": "w12345678", "Student Name": "Alice Smith", "CW1 - Mark": "", "CW1 - Grade": "", "Exam - Mark": "", "Exam - Grade": ""},
            {"Student ID": "12345679/2", "Student Name": "Bob Jones", "CW1 - Mark": "", "CW1 - Grade": "", "Exam - Mark": "", "Exam - Grade": ""}
        ]
        session["converter_mcrf_template"] = base64.b64encode(xls_bytes).decode("utf-8")
        session["converter_mcrf_name"] = "blank_MCRF.xls"
        session.save()

        url = reverse("mcrf_match")
        resp = self.client.post(url, {
            "local_id_col": "Student ID",
            "mcrf_id_col": "Student ID",
            "map_local_0": "CW1",
            "map_mcrf_0": "CW1 - Mark",
            "map_local_1": "Exam",
            "map_mcrf_1": "Exam - Mark"
        })
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp["Content-Type"], "application/vnd.ms-excel")
        self.assertIn("attachment; filename=\"blank_MCRF_Populated.xls\"", resp["Content-Disposition"])
        
        # Verify the populated values and that styles are preserved
        wb = xlrd.open_workbook(file_contents=resp.content, formatting_info=True)
        ws = wb.sheet_by_index(0)
        self.assertEqual(ws.cell_value(4, 2), 85)
        self.assertEqual(ws.cell_value(4, 4), 90)
        self.assertEqual(ws.cell_value(5, 2), 75)
        self.assertEqual(ws.cell_value(5, 4), 80)
        
        # Verify custom bold styling is retained on populated cells
        xf_index = ws.cell_xf_index(4, 2)
        xf = wb.xf_list[xf_index]
        font = wb.font_list[xf.font_index]
        self.assertEqual(font.bold, 1)

    def test_read_uploaded_file_utf16_pseudo_xls(self):
        """Verify read_uploaded_file correctly parses a tab-separated text/CSV file encoded in UTF-16 LE with a BOM (even if named .xls)."""
        from marking_converters.views import read_uploaded_file
        
        # Construct UTF-16 LE tab-separated bytes with LE BOM (b'\xff\xfe')
        text_content = "Student ID\tStudent Name\tCW1\r\n12345678\tAlice\t85\r\n"
        utf16_bytes = b'\xff\xfe' + text_content.encode('utf-16-le')
        
        uploaded_file = SimpleUploadedFile("pseudo_gradebook.xls", utf16_bytes, content_type="application/vnd.ms-excel")
        
        headers, rows = read_uploaded_file(uploaded_file)
        self.assertEqual(headers, ["Student ID", "Student Name", "CW1"])
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["Student ID"], "12345678")
        self.assertEqual(rows[0]["CW1"], "85")

    def test_column_prioritization_in_matching(self):
        """Verify that when both 'Student ID' and 'Username' columns exist, 'Username' is prioritized as the default."""
        session = self.client.session
        session["converter_local_headers"] = ["Student ID", "Username", "CW1"]
        session["converter_local_rows"] = []
        session["converter_bb_headers"] = ["Student ID", "Username", "CW1_Col"]
        session["converter_bb_rows"] = []
        session.save()
        
        # GET request to converter_merge_match
        url = reverse("gradebook_match")
        resp = self.client.get(url)
        
        self.assertEqual(resp.status_code, 200)
        # Verify the context defaults pass 'Username' as the default selected option
        self.assertEqual(resp.context["local_id_default"], "Username")
        self.assertEqual(resp.context["bb_id_default"], "Username")
        
        # Verify the rendered HTML contains the Username option with 'selected' attribute
        # and Student ID does NOT have the selected attribute
        content = resp.content.decode("utf-8")
        self.assertIn('<option value="Username" selected>Username</option>', content)
        self.assertIn('<option value="Student ID" >Student ID</option>', content)



